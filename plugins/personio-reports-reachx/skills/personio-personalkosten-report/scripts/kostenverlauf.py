#!/usr/bin/env python3
"""
REACHX Personio-Report — Arbeitgeberkosten-Zeitverlauf nach Abteilung

Rekonstruiert aus Eintritts- und Austrittsdaten plus dem AKTUELLEN Gehalt einen
monatlichen Arbeitgeberkosten-Verlauf je Abteilung und vergleicht zwei
Jahres-YTD-Zeiträume (Standard: laufendes Jahr gegen Vorjahr, jeweils 1. Januar
bis Stichtag).

WICHTIGE GRENZE — die Personio-API liefert KEINE Gehaltshistorie. Jeder Monat
wird mit dem heutigen Gehalt des Mitarbeiters bewertet. Der Verlauf zeigt damit
den Effekt von Personalauf- und -abbau (wer wann in welcher Abteilung
dazukam/wegging), NICHT den Effekt von Gehaltserhöhungen. Da Vorjahres-Monate
ebenfalls mit dem heutigen Gehalt bewertet werden, ist der ausgewiesene Anstieg
eine UNTERGRENZE des tatsächlichen Anstiegs — Gehaltserhöhungen kommen oben drauf.

Aufruf:
    python3 kostenverlauf.py --employees employees-alle.json --out-dir ./out \\
        --ag-faktor 1.20 --vollzeit-stunden 40 \\
        --von 2024-01 --bis 2026-05 --ytd-stichtag 2026-05-22

Eingabe ist das JSON aus  personio.py employees --include-inactive  (die
ausgeschiedenen Mitarbeiter werden für die Vorjahres-Seite zwingend gebraucht).
"""
from __future__ import annotations

import argparse
import calendar
import datetime as _dt
import json
import sys
from pathlib import Path


# ---------- Hilfen ----------

def to_float(value):
    if value in (None, "", []):
        return None
    try:
        return float(str(value).replace(",", ".").replace(" ", "").strip())
    except (ValueError, TypeError):
        return None


def parse_date(value):
    if not value or not isinstance(value, str):
        return None
    try:
        return _dt.date.fromisoformat(value[:10])
    except ValueError:
        return None


def parse_ym(value: str) -> tuple[int, int]:
    j, m = value.split("-")[:2]
    return int(j), int(m)


def monatsbrutto(emp: dict, vollzeit: float) -> float | None:
    """Aktuelles Monatsbrutto: fix_salary (Intervall beachten), sonst Stundenlohn."""
    fix = to_float(emp.get("fix_salary"))
    if fix and fix > 0:
        roh = str(emp.get("fix_salary_interval") or "").lower()
        return fix / 12.0 if roh in ("yearly", "year", "jaehrlich") else fix
    hourly = to_float(emp.get("hourly_salary"))
    ws = to_float(emp.get("weekly_working_hours"))
    if hourly and hourly > 0 and ws:
        return hourly * ws * 4.33
    return None


def austrittsdatum(emp: dict):
    """Effektives Vertragsende: termination_date, sonst contract_end_date,
    sonst last_working_day."""
    for key in ("termination_date", "contract_end_date", "last_working_day"):
        d = parse_date(emp.get(key))
        if d:
            return d
    return None


def monatsletzter(jahr: int, monat: int) -> _dt.date:
    return _dt.date(jahr, monat, calendar.monthrange(jahr, monat)[1])


def monate_zwischen(von: tuple[int, int], bis: tuple[int, int]):
    """Liefert (jahr, monat)-Tupel von 'von' bis 'bis' inklusive."""
    j, m = von
    while (j, m) <= bis:
        yield j, m
        m += 1
        if m > 12:
            j, m = j + 1, 1


def kosten_im_zeitraum(ag_monat: float, hire, term,
                       start: _dt.date, ende: _dt.date) -> float:
    """Tag-genau anteilige Arbeitgeberkosten eines Mitarbeiters im Zeitraum
    [start, ende]. Jeder Monat zählt anteilig nach beschäftigten Kalendertagen."""
    if ag_monat <= 0:
        return 0.0
    total = 0.0
    start_ym = (start.year, start.month)
    ende_ym = (ende.year, ende.month)
    for jahr, monat in monate_zwischen(start_ym, ende_ym):
        m_erster = _dt.date(jahr, monat, 1)
        m_letzter = monatsletzter(jahr, monat)
        tage_im_monat = m_letzter.day
        seg_start = max(m_erster, start, hire or m_erster)
        seg_ende = min(m_letzter, ende, term or m_letzter)
        if seg_ende >= seg_start:
            tage = (seg_ende - seg_start).days + 1
            total += ag_monat * tage / tage_im_monat
    return total


def beschaeftigt_im_monat(hire, term, jahr: int, monat: int) -> bool:
    f = _dt.date(jahr, monat, 1)
    l = monatsletzter(jahr, monat)
    return (hire is None or hire <= l) and (term is None or term >= f)


# ---------- Kern ----------

def baue_verlauf(emps: list, cfg: dict) -> dict:
    faktor = cfg["ag_faktor"]
    vollzeit = cfg["vollzeit_stunden"]

    # Pro Mitarbeiter: Stammwerte + aktuelle Arbeitgeberkosten pro Monat.
    personen = []
    ohne_gehalt = []
    for e in emps:
        mb = monatsbrutto(e, vollzeit)
        hire = parse_date(e.get("hire_date"))
        term = austrittsdatum(e)
        abt = e.get("department") or "Ohne Abteilung"
        if mb is None or mb <= 0:
            ohne_gehalt.append({"abteilung": abt, "status": e.get("status")})
            continue
        personen.append({
            "abteilung": abt,
            "hire": hire,
            "term": term,
            "ag_monat": mb * faktor,
        })

    abteilungen = sorted({p["abteilung"] for p in personen})

    # Monatsverlauf
    verlauf = []
    for jahr, monat in monate_zwischen(parse_ym(cfg["von"]), parse_ym(cfg["bis"])):
        f = _dt.date(jahr, monat, 1)
        l = monatsletzter(jahr, monat)
        pro_abt = {a: 0.0 for a in abteilungen}
        koepfe = {a: 0 for a in abteilungen}
        for p in personen:
            kosten = kosten_im_zeitraum(p["ag_monat"], p["hire"], p["term"], f, l)
            if kosten > 0:
                pro_abt[p["abteilung"]] += kosten
            if beschaeftigt_im_monat(p["hire"], p["term"], jahr, monat):
                koepfe[p["abteilung"]] += 1
        verlauf.append({
            "monat": f"{jahr}-{monat:02d}",
            "kosten": {a: round(pro_abt[a], 2) for a in abteilungen},
            "kosten_gesamt": round(sum(pro_abt.values()), 2),
            "koepfe": koepfe,
            "koepfe_gesamt": sum(koepfe.values()),
        })

    # YTD-Vergleich: zwei Zeiträume 1. Januar bis Stichtag (gleicher M-D).
    stichtag = _dt.date.fromisoformat(cfg["ytd_stichtag"])
    jahr_neu = stichtag.year
    jahr_alt = jahr_neu - 1
    perioden = {}
    for label, jahr in (("vorjahr", jahr_alt), ("aktuell", jahr_neu)):
        start = _dt.date(jahr, 1, 1)
        ende = _dt.date(jahr, stichtag.month, stichtag.day)
        n_monate = stichtag.month  # angefangene Monate im YTD
        pro_abt = {a: 0.0 for a in abteilungen}
        # durchschnittliche Kopfzahl je Abteilung über die YTD-Monate
        kopf_summe = {a: 0 for a in abteilungen}
        for m in range(1, stichtag.month + 1):
            for p in personen:
                if beschaeftigt_im_monat(p["hire"], p["term"], jahr, m):
                    kopf_summe[p["abteilung"]] += 1
        for p in personen:
            k = kosten_im_zeitraum(p["ag_monat"], p["hire"], p["term"], start, ende)
            if k > 0:
                pro_abt[p["abteilung"]] += k
        perioden[label] = {
            "jahr": jahr,
            "von": start.isoformat(),
            "bis": ende.isoformat(),
            "kosten": {a: round(pro_abt[a], 2) for a in abteilungen},
            "kosten_gesamt": round(sum(pro_abt.values()), 2),
            "koepfe_schnitt": {a: round(kopf_summe[a] / n_monate, 1)
                               for a in abteilungen},
            "koepfe_schnitt_gesamt": round(sum(kopf_summe.values()) / n_monate, 1),
        }

    # Delta je Abteilung
    bridge = []
    pa, pv = perioden["aktuell"], perioden["vorjahr"]
    for a in abteilungen:
        d_kosten = round(pa["kosten"][a] - pv["kosten"][a], 2)
        d_koepfe = round(pa["koepfe_schnitt"][a] - pv["koepfe_schnitt"][a], 1)
        basis = pv["kosten"][a]
        bridge.append({
            "abteilung": a,
            "kosten_vorjahr": pv["kosten"][a],
            "kosten_aktuell": pa["kosten"][a],
            "delta": d_kosten,
            "delta_pct": round(d_kosten / basis * 100, 1) if basis else None,
            "koepfe_vorjahr": pv["koepfe_schnitt"][a],
            "koepfe_aktuell": pa["koepfe_schnitt"][a],
            "delta_koepfe": d_koepfe,
        })
    bridge.sort(key=lambda x: x["delta"], reverse=True)
    delta_gesamt = round(pa["kosten_gesamt"] - pv["kosten_gesamt"], 2)

    return {
        "abteilungen": abteilungen,
        "verlauf": verlauf,
        "ytd": {
            "stichtag": stichtag.isoformat(),
            "vorjahr": pv,
            "aktuell": pa,
            "delta_gesamt": delta_gesamt,
            "delta_pct_gesamt": (round(delta_gesamt / pv["kosten_gesamt"] * 100, 1)
                                 if pv["kosten_gesamt"] else None),
            "bridge": bridge,
        },
        "ohne_gehalt_anzahl": len(ohne_gehalt),
        "ohne_gehalt": ohne_gehalt,
        "personen_mit_gehalt": len(personen),
    }


# ---------- Excel ----------

def schreibe_xlsx(pfad: Path, res: dict, cfg: dict):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rot = PatternFill("solid", fgColor="EC644A")
    grau = PatternFill("solid", fgColor="F2F3F3")
    kopf = Font(name="Calibri", bold=True, color="FFFFFF")
    fett = Font(bold=True)
    EUR = '#,##0 "€"'
    abt = res["abteilungen"]
    wb = Workbook()

    # Blatt 1: Monatsverlauf
    ws = wb.active
    ws.title = "Monatsverlauf"
    headers = ["Monat"] + abt + ["Gesamt"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = rot
        cell.font = kopf
    for r, v in enumerate(res["verlauf"], start=2):
        ws.cell(row=r, column=1, value=v["monat"])
        for c, a in enumerate(abt, start=2):
            ws.cell(row=r, column=c, value=v["kosten"][a]).number_format = EUR
        ws.cell(row=r, column=len(abt) + 2,
                value=v["kosten_gesamt"]).number_format = EUR
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 12
    for i in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18

    # Blatt 2: YTD-Vergleich
    ws = wb.create_sheet("YTD-Vergleich")
    ytd = res["ytd"]
    ja, jv = ytd["aktuell"]["jahr"], ytd["vorjahr"]["jahr"]
    ws["A1"] = (f"YTD-Vergleich Arbeitgeberkosten — "
                f"1.1.–{ytd['stichtag'][8:10]}.{ytd['stichtag'][5:7]}.")
    ws["A1"].font = Font(bold=True, size=12, color="EC644A")
    headers = ["Abteilung", f"Ø Köpfe {jv}", f"AG-Kosten {jv}",
               f"Ø Köpfe {ja}", f"AG-Kosten {ja}", "Delta €", "Delta %",
               "Δ Köpfe"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.fill = rot
        cell.font = kopf
    r = 4
    for b in res["ytd"]["bridge"]:
        ws.cell(row=r, column=1, value=b["abteilung"])
        ws.cell(row=r, column=2, value=b["koepfe_vorjahr"]).number_format = "0.0"
        ws.cell(row=r, column=3, value=b["kosten_vorjahr"]).number_format = EUR
        ws.cell(row=r, column=4, value=b["koepfe_aktuell"]).number_format = "0.0"
        ws.cell(row=r, column=5, value=b["kosten_aktuell"]).number_format = EUR
        ws.cell(row=r, column=6, value=b["delta"]).number_format = EUR
        ws.cell(row=r, column=7,
                value=(b["delta_pct"] / 100 if b["delta_pct"] is not None else None)
                ).number_format = "0.0%"
        ws.cell(row=r, column=8, value=b["delta_koepfe"]).number_format = "+0.0;-0.0"
        r += 1
    for c in range(1, 9):
        ws.cell(row=r, column=c).fill = grau
        ws.cell(row=r, column=c).font = fett
    ws.cell(row=r, column=1, value="Gesamt")
    ws.cell(row=r, column=2,
            value=ytd["vorjahr"]["koepfe_schnitt_gesamt"]).number_format = "0.0"
    ws.cell(row=r, column=3,
            value=ytd["vorjahr"]["kosten_gesamt"]).number_format = EUR
    ws.cell(row=r, column=4,
            value=ytd["aktuell"]["koepfe_schnitt_gesamt"]).number_format = "0.0"
    ws.cell(row=r, column=5,
            value=ytd["aktuell"]["kosten_gesamt"]).number_format = EUR
    ws.cell(row=r, column=6, value=ytd["delta_gesamt"]).number_format = EUR
    ws.cell(row=r, column=7,
            value=(ytd["delta_pct_gesamt"] / 100
                   if ytd["delta_pct_gesamt"] is not None else None)
            ).number_format = "0.0%"
    ws.column_dimensions["A"].width = 22
    for i in range(2, 9):
        ws.column_dimensions[get_column_letter(i)].width = 16

    # Blatt 3: Annahmen
    ws = wb.create_sheet("Annahmen")
    ws["A1"] = "Annahmen & Methodik"
    ws["A1"].font = Font(bold=True, size=12, color="EC644A")
    zeilen = [
        ("Arbeitgeberkosten-Faktor", f"{cfg['ag_faktor']:.2f}"),
        ("Vollzeit-Wochenstunden", f"{cfg['vollzeit_stunden']:g}"),
        ("Verlaufs-Zeitraum", f"{cfg['von']} bis {cfg['bis']}"),
        ("YTD-Stichtag", cfg["ytd_stichtag"]),
        ("Mitarbeiter mit Gehalt", res["personen_mit_gehalt"]),
        ("Mitarbeiter ohne Gehalt (ignoriert)", res["ohne_gehalt_anzahl"]),
    ]
    r = 3
    for label, wert in zeilen:
        ws.cell(row=r, column=1, value=label).font = fett
        ws.cell(row=r, column=2, value=wert)
        r += 1
    r += 1
    hinweise = [
        "GRENZE: Die Personio-API liefert keine Gehaltshistorie. Jeder Monat — "
        "auch im Vorjahr — wird mit dem HEUTIGEN Gehalt bewertet.",
        "Der Verlauf zeigt damit den Effekt von Personalauf-/-abbau, NICHT von "
        "Gehaltserhöhungen.",
        "Da auch Vorjahres-Monate mit dem heutigen (höheren) Gehalt gerechnet "
        "werden, ist der ausgewiesene Anstieg eine UNTERGRENZE — reale "
        "Gehaltserhöhungen erhöhen ihn zusätzlich.",
        "Abteilungs-Zuordnung ist der heutige Stand; frühere Abteilungswechsel "
        "sind nicht abgebildet.",
        "Monatskosten sind tag-genau anteilig (Ein-/Austritt mitten im Monat).",
    ]
    for h in hinweise:
        ws.cell(row=r, column=1, value="• " + h)
        r += 1
    ws.column_dimensions["A"].width = 100

    wb.save(pfad)


# ---------- Hauptlauf ----------

def main(argv=None) -> int:
    p = argparse.ArgumentParser(prog="kostenverlauf.py")
    p.add_argument("--employees", required=True,
                   help="JSON aus  personio.py employees --include-inactive")
    p.add_argument("--out-dir", required=True)
    p.add_argument("--ag-faktor", type=float, default=1.20)
    p.add_argument("--vollzeit-stunden", type=float, default=40.0)
    p.add_argument("--von", default=None, help="Verlaufs-Start JJJJ-MM")
    p.add_argument("--bis", default=None, help="Verlaufs-Ende JJJJ-MM")
    p.add_argument("--ytd-stichtag", default=None,
                   help="YTD-Stichtag JJJJ-MM-TT (Standard: heute)")
    args = p.parse_args(argv)

    src = Path(args.employees).expanduser()
    if not src.exists():
        print(json.dumps({"success": False, "error": f"nicht gefunden: {src}"}),
              file=sys.stderr)
        return 1
    daten = json.loads(src.read_text(encoding="utf-8"))
    emps = daten.get("employees")
    if not emps:
        print(json.dumps({"success": False, "error": "kein 'employees' im JSON"}),
              file=sys.stderr)
        return 1
    if not daten.get("include_inactive"):
        print(json.dumps({"success": False, "error":
              "Datensatz ohne ausgeschiedene Mitarbeiter — bitte mit "
              "'personio.py employees --include-inactive' erzeugen."}),
              file=sys.stderr)
        return 1

    stichtag = args.ytd_stichtag or _dt.date.today().isoformat()
    bis = args.bis or stichtag[:7]
    von = args.von or f"{int(stichtag[:4]) - 2}-01"
    cfg = {
        "ag_faktor": args.ag_faktor,
        "vollzeit_stunden": args.vollzeit_stunden,
        "von": von, "bis": bis, "ytd_stichtag": stichtag,
        "erstellt": _dt.datetime.now().isoformat(timespec="seconds"),
    }

    res = baue_verlauf(emps, cfg)

    out_dir = Path(args.out_dir).expanduser()
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = out_dir / "kostenverlauf.xlsx"
    schreibe_xlsx(xlsx, res, cfg)
    payload = {"success": True, "erstellt_am": cfg["erstellt"],
               "annahmen": cfg, **res}
    json_pfad = out_dir / "kostenverlauf-daten.json"
    json_pfad.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
                         encoding="utf-8")

    ytd = res["ytd"]
    print(json.dumps({
        "success": True,
        "geschrieben": [str(xlsx), str(json_pfad)],
        "ytd_vorjahr": ytd["vorjahr"]["kosten_gesamt"],
        "ytd_aktuell": ytd["aktuell"]["kosten_gesamt"],
        "delta_gesamt": ytd["delta_gesamt"],
        "delta_pct_gesamt": ytd["delta_pct_gesamt"],
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
