#!/usr/bin/env python3
"""
REACHX Personio-Report — Aggregation und Excel-Mappe

Liest das von  personio.py employees  erzeugte JSON, berechnet die
Headcount-, FTE-, Personalkosten- und Arbeitgeberkosten-Auswertung und
schreibt zwei Dateien ins Ausgabe-Verzeichnis:

  - personalkosten-report.xlsx   formatierte Excel-Mappe (vier Blaetter)
  - report-daten.json            verdichtete Kennzahlen fuer den HTML-Report

ARBEITGEBERKOSTEN sind eine SCHAETZUNG: Die Personio-API liefert das
vertragliche Bruttogehalt, nicht die tatsaechlichen Arbeitgeberkosten. Der
Skill multipliziert das Brutto mit einem konfigurierbaren AG-Faktor
(--ag-faktor, Standard 1.20 fuer ~20 % Arbeitgeberanteil zur
Sozialversicherung in Deutschland). Der Faktor steht in jedem Output.

Aufruf (Beispiel):
    python3 build-report.py --employees employees.json --out-dir ./report \\
        --ag-faktor 1.20 --vollzeit-stunden 40 \\
        --gehalt-attribut fix_salary --gehalt-intervall auto \\
        --bonus-attribut bonus --bonus-intervall jaehrlich

Mit --anonym werden Klarnamen durch Codes (MA-001 ...) ersetzt.
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import sys
from pathlib import Path

WOCHEN_PRO_MONAT = 4.33        # Umrechnung Stundenlohn -> Monatsbrutto


# ---------- Hilfen ----------

def to_float(value) -> float | None:
    if value in (None, "", []):
        return None
    try:
        return float(str(value).replace(",", ".").replace(" ", "").strip())
    except (ValueError, TypeError):
        return None


def parse_date(value) -> _dt.date | None:
    if not value or not isinstance(value, str):
        return None
    try:
        return _dt.date.fromisoformat(value[:10])
    except ValueError:
        return None


def fmt_name(emp: dict) -> str:
    parts = [str(emp.get("first_name") or "").strip(),
             str(emp.get("last_name") or "").strip()]
    name = " ".join(p for p in parts if p)
    return name or f"ID {emp.get('id', '?')}"


def _monatsbrutto_aus(rohwert, intervall: str, wochenstunden) -> float | None:
    """Rechnet einen Gehaltswert auf ein Monatsbrutto um."""
    if not rohwert or rohwert <= 0:
        return None
    if intervall == "stunde":
        return rohwert * wochenstunden * WOCHEN_PRO_MONAT if wochenstunden else None
    if intervall == "jaehrlich":
        return rohwert / 12.0
    return rohwert        # monatlich


def _fix_salary_intervall(emp: dict) -> str:
    roh = str(emp.get("fix_salary_interval") or "").lower()
    return "jaehrlich" if roh in ("yearly", "year", "jaehrlich") else "monatlich"


# ---------- Kern-Berechnung ----------

def berechne_mitarbeiter(emp: dict, cfg: dict) -> dict:
    """Leitet pro Mitarbeiter Wochenstunden, FTE, Brutto und AG-Kosten ab."""
    wochenstunden = to_float(emp.get("weekly_working_hours"))
    vollzeit = cfg["vollzeit_stunden"]
    fte = round(wochenstunden / vollzeit, 3) if wochenstunden and vollzeit else None

    attr = cfg["gehalt_attribut"]
    monatsbrutto = None
    gehalt_basis = None
    if attr == "auto":
        # Festgehalt bevorzugen, sonst auf Stundenlohn ausweichen.
        fix = to_float(emp.get("fix_salary"))
        if fix and fix > 0:
            monatsbrutto = _monatsbrutto_aus(fix, _fix_salary_intervall(emp),
                                             wochenstunden)
            gehalt_basis = "fix_salary"
        else:
            hourly = to_float(emp.get("hourly_salary"))
            if hourly and hourly > 0:
                monatsbrutto = _monatsbrutto_aus(hourly, "stunde", wochenstunden)
                gehalt_basis = "hourly_salary"
    else:
        rohwert = to_float(emp.get(attr))
        intervall = cfg["gehalt_intervall"]
        if intervall == "auto":
            if attr == "hourly_salary":
                intervall = "stunde"
            elif attr == "fix_salary":
                intervall = _fix_salary_intervall(emp)
            else:
                intervall = "monatlich"
        monatsbrutto = _monatsbrutto_aus(rohwert, intervall, wochenstunden)
        if monatsbrutto is not None:
            gehalt_basis = attr

    jahresbrutto = monatsbrutto * 12.0 if monatsbrutto is not None else None

    bonus_jahr = 0.0
    if cfg["bonus_attribut"]:
        bonus_roh = to_float(emp.get(cfg["bonus_attribut"]))
        if bonus_roh and bonus_roh > 0:
            bonus_jahr = (bonus_roh * 12.0 if cfg["bonus_intervall"] == "monatlich"
                          else bonus_roh)

    jahresbrutto_gesamt = None
    if jahresbrutto is not None:
        jahresbrutto_gesamt = jahresbrutto + bonus_jahr

    ag_kosten_jahr = None
    if jahresbrutto_gesamt is not None:
        ag_kosten_jahr = jahresbrutto_gesamt * cfg["ag_faktor"]

    return {
        "id": emp.get("id"),
        "name": emp.get("_anzeigename"),
        "abteilung": emp.get("department") or "Ohne Abteilung",
        "status": str(emp.get("status") or "").lower() or "unbekannt",
        "position": emp.get("position") or "",
        "eintritt": emp.get("hire_date"),
        "austritt": emp.get("termination_date") or emp.get("contract_end_date"),
        "wochenstunden": wochenstunden,
        "fte": fte,
        "monatsbrutto": monatsbrutto,
        "jahresbrutto": jahresbrutto,
        "bonus_jahr": bonus_jahr,
        "jahresbrutto_gesamt": jahresbrutto_gesamt,
        "ag_kosten_jahr": ag_kosten_jahr,
        "ag_kosten_monat": ag_kosten_jahr / 12.0 if ag_kosten_jahr is not None else None,
        "gehalt_basis": gehalt_basis,
        "hat_gehalt": jahresbrutto_gesamt is not None,
    }


def aggregiere(zeilen: list, cfg: dict) -> dict:
    stichtag = cfg["stichtag"]
    vor_12m = stichtag - _dt.timedelta(days=365)

    headcount = len(zeilen)
    fte_summe = round(sum(z["fte"] for z in zeilen if z["fte"]), 2)
    mit_gehalt = [z for z in zeilen if z["hat_gehalt"]]
    ohne_gehalt = [z for z in zeilen if not z["hat_gehalt"]]

    brutto_jahr = sum(z["jahresbrutto_gesamt"] for z in mit_gehalt)
    ag_jahr = sum(z["ag_kosten_jahr"] for z in mit_gehalt)
    n = len(mit_gehalt)
    fte_mit_gehalt = sum(z["fte"] for z in mit_gehalt if z["fte"]) or 0

    # Abteilungen
    abt: dict = {}
    for z in zeilen:
        a = abt.setdefault(z["abteilung"], {
            "abteilung": z["abteilung"], "headcount": 0, "fte": 0.0,
            "mit_gehalt": 0, "brutto_jahr": 0.0, "ag_jahr": 0.0})
        a["headcount"] += 1
        a["fte"] += z["fte"] or 0
        if z["hat_gehalt"]:
            a["mit_gehalt"] += 1
            a["brutto_jahr"] += z["jahresbrutto_gesamt"]
            a["ag_jahr"] += z["ag_kosten_jahr"]
    abteilungen = sorted(abt.values(), key=lambda x: x["ag_jahr"], reverse=True)
    for a in abteilungen:
        a["fte"] = round(a["fte"], 2)

    # Status-Verteilung
    status_verteilung: dict = {}
    for z in zeilen:
        status_verteilung[z["status"]] = status_verteilung.get(z["status"], 0) + 1

    # Gehalts-Basis-Verteilung (Festgehalt vs. Stundenlohn vs. Custom)
    basis_verteilung: dict = {}
    for z in mit_gehalt:
        b = z["gehalt_basis"] or "unbekannt"
        basis_verteilung[b] = basis_verteilung.get(b, 0) + 1

    # Fluktuation
    eintritte = [z for z in zeilen
                 if (d := parse_date(z["eintritt"])) and vor_12m <= d <= stichtag]
    austritte = [z for z in zeilen
                 if (d := parse_date(z["austritt"])) and vor_12m <= d <= stichtag]
    fluktuationsrate = round(len(austritte) / headcount, 4) if headcount else None

    # Betriebszugehoerigkeit
    jahre = [(stichtag - d).days / 365.25 for z in zeilen
             if (d := parse_date(z["eintritt"]))]
    betriebszugehoerigkeit = round(sum(jahre) / len(jahre), 1) if jahre else None

    return {
        "headcount": headcount,
        "fte_summe": fte_summe,
        "mit_gehalt": n,
        "ohne_gehalt": len(ohne_gehalt),
        "ohne_gehalt_namen": [z["name"] for z in ohne_gehalt],
        "kostendaten_verfuegbar": n > 0,
        "personalkosten": {
            "brutto_jahr_summe": round(brutto_jahr, 2),
            "ag_kosten_jahr_summe": round(ag_jahr, 2),
            "ag_kosten_monat_summe": round(ag_jahr / 12.0, 2),
            "brutto_jahr_schnitt": round(brutto_jahr / n, 2) if n else None,
            "ag_kosten_jahr_schnitt_kopf": round(ag_jahr / n, 2) if n else None,
            "ag_kosten_jahr_je_fte": round(ag_jahr / fte_mit_gehalt, 2)
                                     if fte_mit_gehalt else None,
        },
        "abteilungen": abteilungen,
        "status_verteilung": status_verteilung,
        "gehalt_basis_verteilung": basis_verteilung,
        "fluktuation": {
            "eintritte_12m": len(eintritte),
            "austritte_12m": len(austritte),
            "fluktuationsrate": fluktuationsrate,
            "betriebszugehoerigkeit_schnitt_jahre": betriebszugehoerigkeit,
        },
    }


# ---------- Excel ----------

def schreibe_xlsx(pfad: Path, zeilen: list, agg: dict, cfg: dict, warnungen: list):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise SystemExit(
            "openpyxl fehlt. Installieren mit:\n"
            "  pip3 install -r requirements-personio.txt")

    rot = PatternFill("solid", fgColor="EC644A")
    grau = PatternFill("solid", fgColor="F2F3F3")
    kopf_font = Font(name="Calibri", bold=True, color="FFFFFF")
    fett = Font(bold=True)
    EUR = '#,##0 "€"'
    rechts = Alignment(horizontal="right")

    wb = Workbook()

    def kopfzeile(ws, headers, row=1):
        for c, text in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=c, value=text)
            cell.fill = rot
            cell.font = kopf_font
        ws.freeze_panes = ws.cell(row=row + 1, column=1)

    def breiten(ws, widths):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # --- Blatt 1: Uebersicht ---
    ws = wb.active
    ws.title = "Übersicht"
    pk = agg["personalkosten"]
    fl = agg["fluktuation"]
    ws["A1"] = "REACHX · Personalkosten-Report"
    ws["A1"].font = Font(bold=True, size=14, color="EC644A")
    ws["A2"] = f"Stichtag {cfg['stichtag'].isoformat()} · erstellt {cfg['erstellt']}"
    ws["A2"].font = Font(italic=True, color="999DA1")

    kpis = [
        ("Kennzahl", "Wert", None),
        ("Mitarbeiter (Headcount)", agg["headcount"], "0"),
        ("Vollzeitäquivalente (FTE)", agg["fte_summe"], "0.0"),
        ("Mitarbeiter mit Gehaltsangabe", agg["mit_gehalt"], "0"),
        ("Bruttogehälter p.a. (gesamt)", pk["brutto_jahr_summe"], EUR),
        ("Arbeitgeberkosten p.a. (gesamt, geschätzt)",
         pk["ag_kosten_jahr_summe"], EUR),
        ("Arbeitgeberkosten p.M. (gesamt, geschätzt)",
         pk["ag_kosten_monat_summe"], EUR),
        ("Ø Bruttogehalt p.a. je Kopf", pk["brutto_jahr_schnitt"], EUR),
        ("Ø Arbeitgeberkosten p.a. je Kopf",
         pk["ag_kosten_jahr_schnitt_kopf"], EUR),
        ("Ø Arbeitgeberkosten p.a. je FTE", pk["ag_kosten_jahr_je_fte"], EUR),
        ("Eintritte (letzte 12 Monate)", fl["eintritte_12m"], "0"),
        ("Austritte (letzte 12 Monate)", fl["austritte_12m"], "0"),
        ("Fluktuationsrate",
         fl["fluktuationsrate"], "0.0%"),
        ("Ø Betriebszugehörigkeit (Jahre)",
         fl["betriebszugehoerigkeit_schnitt_jahre"], "0.0"),
    ]
    start = 4
    for r, (label, wert, nf) in enumerate(kpis, start=start):
        a = ws.cell(row=r, column=1, value=label)
        b = ws.cell(row=r, column=2, value=wert)
        if r == start:
            a.fill = rot; a.font = kopf_font
            b.fill = rot; b.font = kopf_font
        else:
            a.font = fett
            b.alignment = rechts
            if nf and wert is not None:
                b.number_format = nf
    breiten(ws, [42, 22])
    ws.freeze_panes = "A5"

    # --- Blatt 2: Mitarbeiter ---
    ws = wb.create_sheet("Mitarbeiter")
    headers = ["Name", "Abteilung", "Position", "Status", "Eintritt",
               "Austritt", "Wochenstunden", "FTE", "Monatsbrutto",
               "Jahresbrutto inkl. Bonus", "Arbeitgeberkosten p.a. (geschätzt)"]
    kopfzeile(ws, headers)
    for r, z in enumerate(zeilen, start=2):
        ws.cell(row=r, column=1, value=z["name"])
        ws.cell(row=r, column=2, value=z["abteilung"])
        ws.cell(row=r, column=3, value=z["position"])
        ws.cell(row=r, column=4, value=z["status"])
        ws.cell(row=r, column=5, value=z["eintritt"] or "")
        ws.cell(row=r, column=6, value=z["austritt"] or "")
        ws.cell(row=r, column=7, value=z["wochenstunden"])
        c8 = ws.cell(row=r, column=8, value=z["fte"])
        c8.number_format = "0.00"
        for col, key in ((9, "monatsbrutto"), (10, "jahresbrutto_gesamt"),
                         (11, "ag_kosten_jahr")):
            cell = ws.cell(row=r, column=col, value=z[key])
            cell.number_format = EUR
    breiten(ws, [26, 20, 22, 12, 12, 12, 14, 8, 16, 22, 26])

    # --- Blatt 3: Abteilungen ---
    ws = wb.create_sheet("Abteilungen")
    headers = ["Abteilung", "Headcount", "FTE", "Mit Gehalt",
               "Bruttogehälter p.a.", "Arbeitgeberkosten p.a. (geschätzt)"]
    kopfzeile(ws, headers)
    r = 2
    for a in agg["abteilungen"]:
        ws.cell(row=r, column=1, value=a["abteilung"])
        ws.cell(row=r, column=2, value=a["headcount"])
        ws.cell(row=r, column=3, value=round(a["fte"], 2)).number_format = "0.00"
        ws.cell(row=r, column=4, value=a["mit_gehalt"])
        ws.cell(row=r, column=5, value=round(a["brutto_jahr"], 2)).number_format = EUR
        ws.cell(row=r, column=6, value=round(a["ag_jahr"], 2)).number_format = EUR
        r += 1
    # Summenzeile
    for c in range(1, 7):
        ws.cell(row=r, column=c).fill = grau
        ws.cell(row=r, column=c).font = fett
    ws.cell(row=r, column=1, value="Gesamt")
    ws.cell(row=r, column=2, value=agg["headcount"])
    ws.cell(row=r, column=3, value=agg["fte_summe"]).number_format = "0.00"
    ws.cell(row=r, column=4, value=agg["mit_gehalt"])
    ws.cell(row=r, column=5,
            value=pk["brutto_jahr_summe"]).number_format = EUR
    ws.cell(row=r, column=6,
            value=pk["ag_kosten_jahr_summe"]).number_format = EUR
    breiten(ws, [24, 12, 10, 12, 22, 26])

    # --- Blatt 4: Annahmen & Hinweise ---
    ws = wb.create_sheet("Annahmen & Hinweise")
    ws["A1"] = "Annahmen & Hinweise"
    ws["A1"].font = Font(bold=True, size=12, color="EC644A")
    annahmen = [
        ("Arbeitgeberkosten-Faktor", f"{cfg['ag_faktor']:.2f}"),
        ("Vollzeit-Wochenstunden (FTE-Basis)", f"{cfg['vollzeit_stunden']:g}"),
        ("Gehalts-Attribut", cfg["gehalt_attribut"]),
        ("Gehalts-Intervall", cfg["gehalt_intervall"]),
        ("Bonus-Attribut", cfg["bonus_attribut"] or "— (nicht berücksichtigt)"),
        ("Bonus-Intervall", cfg["bonus_intervall"]),
        ("Stichtag", cfg["stichtag"].isoformat()),
        ("Namen anonymisiert", "ja" if cfg["anonym"] else "nein"),
    ]
    r = 3
    for label, wert in annahmen:
        ws.cell(row=r, column=1, value=label).font = fett
        ws.cell(row=r, column=2, value=wert)
        r += 1
    r += 1
    ws.cell(row=r, column=1,
            value="WICHTIG: Die Arbeitgeberkosten sind eine Schätzung "
                  "(Bruttogehalt × AG-Faktor), keine Ist-Abrechnungswerte.").font = \
        Font(italic=True, color="BF4C4A")
    r += 2
    if warnungen:
        ws.cell(row=r, column=1, value="Datenlücken / Warnungen:").font = fett
        r += 1
        for w in warnungen:
            ws.cell(row=r, column=1, value=f"• {w}")
            r += 1
    breiten(ws, [40, 50])

    wb.save(pfad)


# ---------- Hauptlauf ----------

def main(argv=None) -> int:
    p = argparse.ArgumentParser(prog="build-report.py",
                                description="Personio-Daten aggregieren und "
                                            "Excel-Mappe schreiben.")
    p.add_argument("--employees", required=True,
                   help="JSON-Datei aus  personio.py employees.")
    p.add_argument("--out-dir", required=True, help="Ausgabe-Verzeichnis.")
    p.add_argument("--ag-faktor", type=float, default=1.20,
                   help="Arbeitgeberkosten-Faktor (Standard 1.20).")
    p.add_argument("--vollzeit-stunden", type=float, default=40.0,
                   help="Wochenstunden einer Vollzeitstelle (FTE-Basis).")
    p.add_argument("--gehalt-attribut", default="auto",
                   help="Gehalts-Attribut. 'auto' (Standard) nimmt fix_salary, "
                        "sonst hilfsweise hourly_salary. Alternativ ein fester "
                        "Schlüssel, z.B. ein Custom-Attribut dynamic_NNNN.")
    p.add_argument("--gehalt-intervall", default="auto",
                   choices=["auto", "monatlich", "jaehrlich", "stunde"],
                   help="Intervall des Gehalts-Attributs.")
    p.add_argument("--bonus-attribut", default="bonus",
                   help="Attribut-Schlüssel für Bonus; leer = ignorieren.")
    p.add_argument("--bonus-intervall", default="jaehrlich",
                   choices=["jaehrlich", "monatlich"])
    p.add_argument("--stichtag", default=None,
                   help="Stichtag YYYY-MM-DD (Standard: heute).")
    p.add_argument("--anonym", action="store_true",
                   help="Klarnamen durch Codes MA-001 … ersetzen.")
    args = p.parse_args(argv)

    src = Path(args.employees).expanduser()
    if not src.exists():
        print(json.dumps({"success": False,
                          "error": f"Datei nicht gefunden: {src}"}), file=sys.stderr)
        return 1
    daten = json.loads(src.read_text(encoding="utf-8"))
    employees = daten.get("employees")
    if employees is None:
        print(json.dumps({"success": False,
                          "error": "JSON enthält kein 'employees' — wurde "
                                   "personio.py employees verwendet?"}),
              file=sys.stderr)
        return 1

    stichtag = (_dt.date.fromisoformat(args.stichtag) if args.stichtag
                else _dt.date.today())
    cfg = {
        "ag_faktor": args.ag_faktor,
        "vollzeit_stunden": args.vollzeit_stunden,
        "gehalt_attribut": args.gehalt_attribut,
        "gehalt_intervall": args.gehalt_intervall,
        "bonus_attribut": args.bonus_attribut.strip(),
        "bonus_intervall": args.bonus_intervall,
        "stichtag": stichtag,
        "anonym": args.anonym,
        "erstellt": _dt.datetime.now().isoformat(timespec="seconds"),
    }

    # Anzeigenamen festlegen (ggf. anonymisiert)
    employees = sorted(employees, key=lambda e: e.get("id") or 0)
    for i, emp in enumerate(employees, start=1):
        emp["_anzeigename"] = (f"MA-{i:03d}" if args.anonym else fmt_name(emp))

    zeilen = [berechne_mitarbeiter(e, cfg) for e in employees]
    zeilen.sort(key=lambda z: (z["ag_kosten_jahr"] or -1), reverse=True)
    agg = aggregiere(zeilen, cfg)

    # Warnungen
    warnungen: list = []
    if agg["mit_gehalt"] == 0:
        quelle = ("den Standard-Feldern fix_salary / hourly_salary"
                  if cfg["gehalt_attribut"] == "auto"
                  else f"dem Attribut '{cfg['gehalt_attribut']}'")
        warnungen.append(
            f"Kein einziger Mitarbeiter hat einen Gehaltswert in {quelle}. "
            f"Vermutlich fehlt dem Personio-API-Zugang die Attribut-Freigabe "
            f"für Gehaltsfelder, oder das Gehalt liegt in einem Custom-Attribut. "
            f"Prüfen mit:  personio.py attributes")
    elif agg["ohne_gehalt"] > 0:
        warnungen.append(
            f"{agg['ohne_gehalt']} von {agg['headcount']} Mitarbeitern ohne "
            f"Gehaltsangabe — sie fehlen in allen Kostensummen.")
    if daten.get("include_inactive"):
        warnungen.append("Datensatz enthält auch inaktive (ausgeschiedene) "
                         "Mitarbeiter — Kennzahlen entsprechend lesen.")

    out_dir = Path(args.out_dir).expanduser()
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx_pfad = out_dir / "personalkosten-report.xlsx"
    schreibe_xlsx(xlsx_pfad, zeilen, agg, cfg, warnungen)

    report_daten = {
        "success": True,
        "erstellt_am": cfg["erstellt"],
        "stichtag": stichtag.isoformat(),
        "quelle": "Personio REST API v1 (First-Party, read-only)",
        "annahmen": {
            "ag_faktor": cfg["ag_faktor"],
            "vollzeit_stunden": cfg["vollzeit_stunden"],
            "gehalt_attribut": cfg["gehalt_attribut"],
            "gehalt_intervall": cfg["gehalt_intervall"],
            "bonus_attribut": cfg["bonus_attribut"],
            "bonus_intervall": cfg["bonus_intervall"],
            "anonymisiert": cfg["anonym"],
        },
        "kennzahlen": {
            "headcount": agg["headcount"],
            "fte_summe": agg["fte_summe"],
            "mit_gehalt": agg["mit_gehalt"],
            "ohne_gehalt": agg["ohne_gehalt"],
            "kostendaten_verfuegbar": agg["kostendaten_verfuegbar"],
            "personalkosten": agg["personalkosten"],
            "fluktuation": agg["fluktuation"],
            "status_verteilung": agg["status_verteilung"],
            "gehalt_basis_verteilung": agg["gehalt_basis_verteilung"],
        },
        "abteilungen": agg["abteilungen"],
        "ohne_gehalt_namen": agg["ohne_gehalt_namen"],
        "warnungen": warnungen,
        "dateien": {"xlsx": xlsx_pfad.name},
    }
    json_pfad = out_dir / "report-daten.json"
    json_pfad.write_text(json.dumps(report_daten, ensure_ascii=False, indent=2)
                         + "\n", encoding="utf-8")

    print(json.dumps({
        "success": True,
        "geschrieben": [str(xlsx_pfad), str(json_pfad)],
        "headcount": agg["headcount"],
        "mit_gehalt": agg["mit_gehalt"],
        "ag_kosten_jahr_summe": agg["personalkosten"]["ag_kosten_jahr_summe"],
        "warnungen": warnungen,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
